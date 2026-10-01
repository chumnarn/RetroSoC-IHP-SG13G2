#!/usr/bin/env python3
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("RetroSoC-Full-Chip-Manual-TH.md")
OUTPUT = Path("RetroSoC-Full-Chip-Manual-TH.docx")
BODY_FONT = "Sarabun"
MONO_FONT = "DejaVu Sans Mono"


def set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    language = rpr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        rpr.append(language)
    language.set(qn("w:val"), "th-TH")
    language.set(qn("w:eastAsia"), "th-TH")
    language.set(qn("w:bidi"), "th-TH")
    if size is not None:
        size_cs = rpr.find(qn("w:szCs"))
        if size_cs is None:
            size_cs = OxmlElement("w:szCs")
            rpr.append(size_cs)
        size_cs.set(qn("w:val"), str(int(size * 2)))


def set_style_font(style, name: str) -> None:
    style.font.name = name
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    language = rpr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        rpr.append(language)
    language.set(qn("w:val"), "th-TH")
    language.set(qn("w:eastAsia"), "th-TH")
    language.set(qn("w:bidi"), "th-TH")


def set_cell_border(cell, color: str = "B7B7B7", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def no_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def add_inline(paragraph, text: str, *, mono: bool = False) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(clean_inline(text[cursor : match.start()]))
            set_font(run, BODY_FONT)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, BODY_FONT, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, MONO_FONT, 8.2)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(clean_inline(text[cursor:]))
        set_font(run, MONO_FONT if mono else BODY_FONT, 8.2 if mono else None)


def add_code(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["No Spacing"]
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line + ("\n" if index < len(lines) - 1 else ""))
        set_font(run, MONO_FONT, 7.4)


def add_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        no_row_split(row)
        if row_index == 0:
            repeat_header(row)
        for col_index in range(columns):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            value = values[col_index] if col_index < len(values) else ""
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_font(run, BODY_FONT, 8.2, bold=(row_index == 0))
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_markdown(document: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    code: list[str] | None = None
    first_h1 = True
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if code is None:
                code = []
            else:
                add_code(document, code)
                code = None
            index += 1
            continue
        if code is not None:
            code.append(line)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?[ :|-]+\|", lines[index + 1]):
            rows: list[list[str]] = []
            rows.append([item.strip() for item in line.strip().strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([item.strip() for item in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            title = clean_inline(heading.group(2))
            if level == 1 and first_h1:
                paragraph = document.add_paragraph(style="Title")
                add_inline(paragraph, title)
                first_h1 = False
            else:
                paragraph = document.add_paragraph(style=f"Heading {level}")
                add_inline(paragraph, title)
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet or numbered:
            if bullet:
                paragraph = document.add_paragraph(style="List Bullet")
                value = bullet.group(1)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.24)
                paragraph.paragraph_format.first_line_indent = Inches(-0.24)
                value = line.strip()
            add_inline(paragraph, value)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        parts = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            if not candidate.strip() or candidate.startswith(("#", "```", "|", "- ", "* ")) or re.match(r"^\d+\.\s+", candidate):
                break
            parts.append(candidate.strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(parts))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("หน้า ")
    set_font(run, BODY_FONT, 8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    for style_name in ("Normal", "Body Text", "List Bullet", "List Number"):
        style = document.styles[style_name]
        set_style_font(style, BODY_FONT)
        style.font.size = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08
    title = document.styles["Title"]
    set_style_font(title, BODY_FONT)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for level, size in ((1, 16), (2, 13), (3, 11)):
        style = document.styles[f"Heading {level}"]
        set_style_font(style, BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
    parse_markdown(document, SOURCE.read_text(encoding="utf-8"))
    footer = section.footer
    add_page_number(footer.paragraphs[0])
    document.core_properties.title = "คู่มือ RetroSoC Full Chip Implementation using LibreLane and IHP SG13G2"
    document.core_properties.subject = "ขั้นตอนรัน full chip และ validation gates"
    document.core_properties.author = "OpenAI"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
